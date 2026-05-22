from fastapi import FastAPI, Depends, Request, HTTPException, status, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, StreamingResponse
from fastapi.security import HTTPAuthorizationCredentials, HTTPBearer
from pydantic import BaseModel
from typing import Optional, List, Dict, Any
from datetime import datetime, timedelta, date
from jose import jwt, JWTError
from dotenv import load_dotenv
import os
import smtplib
from email.mime.text import MIMEText
from email.header import Header
import asyncio
import uuid
import requests
import json
import time
import re
import aiohttp
import aiofiles
import PyPDF2
from docx import Document
import openpyxl
from langdetect import detect
from fastapi.responses import StreamingResponse
from app.utils.logger import log_error, log_info, log_warning
from app.utils.telegram import send_telegram_message_sync

# ===== Конфиг =====
from app.config import config

# ===== БД =====
from sqlalchemy.orm import Session
from app.database import get_db, init_db
from app.models import User, KnowledgeBase

# ===== Security =====
from app.utils.security import (
    verify_password, 
    get_current_user, 
    create_access_token, 
    oauth2_scheme,
    decode_access_token
)

# ===== Redis =====
from app.utils.redis_client import (
    get_session, add_to_memory, reset_memory, build_history_text,
    get_cached_response, set_cached_response
)

# ===== Балансировка моделей =====
from app.utils.model_balancer import select_model, record_model_result

# ===== Celery =====
from app.tasks import celery_app, generate_response_task, get_task_result, save_task_result

# ===== Контроль стоимости =====
from app.utils.cost_manager import (
    calculate_cost, check_budget, update_user_cost, 
    get_remaining_budget_percent, get_model_by_budget
)

# ===== ROUTES =====
from app.routes.users import router as users_router
from app.routes.admin import router as admin_router
from app.routes.password_reset import router as password_reset_router
from app.routes.scripts import router as scripts_router
from app.routes.knowledge import router as knowledge_router
from app.routes.bot import router as bot_router
from app.routes.password_reset import router as password_reset_router

# ===== PROMPTS =====
from app.prompts import build_generate_prompt, build_stream_prompt, build_improve_prompt

# ===== Groq =====
from groq import AsyncGroq

# ===== Tavily Search =====
from app.utils.tavily_search import search_legislation, should_search_external

# ===== ENV =====
load_dotenv()
EMAIL_ADDRESS = config.EMAIL_ADDRESS
EMAIL_PASSWORD = config.EMAIL_PASSWORD
CORS_ORIGINS = config.CORS_ORIGINS
GROQ_API_KEY = config.GROQ_API_KEY

# ===== FastAPI =====
app = FastAPI(title="ReplyFlow AI", version="1.0.0")

security = HTTPBearer(auto_error=False)

# ===== CORS =====
app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "https://classy-vacherin-53dad1.netlify.app",
        "http://localhost:3000"
    ],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# Папка для загрузок
UPLOAD_DIR = "uploads/knowledge"
os.makedirs(UPLOAD_DIR, exist_ok=True)

# ===== ФУНКЦИИ ДЛЯ ИЗВЛЕЧЕНИЯ ТЕКСТА ИЗ ФАЙЛОВ =====
def extract_text_from_file(file_path: str, file_type: str) -> str:
    try:
        if file_type == "pdf":
            with open(file_path, "rb") as f:
                reader = PyPDF2.PdfReader(f)
                text = "\n".join([page.extract_text() for page in reader.pages if page.extract_text()])
            return text[:5000] if text else "PDF не содержит текста"
        elif file_type == "docx":
            doc = Document(file_path)
            text = "\n".join([para.text for para in doc.paragraphs if para.text])
            return text[:5000]
        elif file_type == "xlsx":
            wb = openpyxl.load_workbook(file_path, data_only=True)
            text = ""
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    text += " ".join([str(cell) for cell in row if cell]) + "\n"
            return text[:5000]
        elif file_type == "txt":
            with open(file_path, "r", encoding="utf-8") as f:
                return f.read()[:5000]
        else:
            return f"Неподдерживаемый тип: {file_type}"
    except Exception as e:
        return f"Ошибка извлечения: {e}"
        
# ===== ФУНКЦИИ ДЛЯ CRM =====
async def verify_hubspot_api_key(api_key: str) -> bool:
    try:
        async with aiohttp.ClientSession() as session:
            async with session.get(
                "https://api.hubapi.com/crm/v3/objects/contacts",
                headers={"Authorization": f"Bearer {api_key}"}
            ) as resp:
                return resp.status == 200
    except:
        return False


async def verify_amocrm_api_key(api_key: str, subdomain: str) -> bool:
    try:
        async with aiohttp.ClientSession() as session:
            async with session.get(
                f"https://{subdomain}.amocrm.ru/api/v4/account",
                headers={"Authorization": f"Bearer {api_key}"}
            ) as resp:
                return resp.status == 200
    except:
        return False


async def verify_pipedrive_api_key(api_key: str) -> bool:
    try:
        async with aiohttp.ClientSession() as session:
            async with session.get(
                f"https://api.pipedrive.com/v1/users?api_token={api_key}"
            ) as resp:
                return resp.status == 200
    except:
        return False


async def send_to_hubspot(api_key: str, lead_id: str, question: str, answer: str):
    try:
        async with aiohttp.ClientSession() as session:
            payload = {
                "engagement": {"active": True, "type": "NOTE"},
                "associations": {"dealIds": [int(lead_id)] if lead_id.isdigit() else []},
                "metadata": {"body": f"🤖 ReplyFlow AI\n\n📝 Вопрос:\n{question}\n\n💡 Ответ:\n{answer}"}
            }
            async with session.post(
                "https://api.hubapi.com/engagements/v1/engagements",
                headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"},
                json=payload
            ) as resp:
                return resp.status in [200, 201]
    except Exception as e:
        print(f"[CRM] HubSpot error: {e}")
        return False


async def send_to_amocrm(api_key: str, subdomain: str, lead_id: str, question: str, answer: str):
    try:
        async with aiohttp.ClientSession() as session:
            payload = [{
                "element_id": int(lead_id),
                "element_type": 2,
                "note_type": 4,
                "text": f"🤖 ReplyFlow AI\n\n📝 Вопрос:\n{question}\n\n💡 Ответ:\n{answer}"
            }]
            async with session.post(
                f"https://{subdomain}.amocrm.ru/api/v4/notes",
                headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"},
                json=payload
            ) as resp:
                return resp.status in [200, 201]
    except Exception as e:
        print(f"[CRM] amoCRM error: {e}")
        return False


async def send_to_pipedrive(api_key: str, lead_id: str, question: str, answer: str):
    try:
        async with aiohttp.ClientSession() as session:
            payload = {"content": f"🤖 ReplyFlow AI\n\n📝 Вопрос:\n{question}\n\n💡 Ответ:\n{answer}"}
            async with session.post(
                f"https://api.pipedrive.com/v1/deals/{lead_id}/activities?api_token={api_key}",
                json=payload
            ) as resp:
                return resp.status in [200, 201]
    except Exception as e:
        print(f"[CRM] Pipedrive error: {e}")
        return False


async def send_to_bitrix24(webhook_url: str, lead_id: str, question: str, answer: str):
    if not webhook_url or not lead_id:
        return False
    
    api_method = "crm.timeline.comment.add"
    full_url = f"{webhook_url.rstrip('/')}/{api_method}.json"
    
    payload = {
        "fields": {
            "ENTITY_ID": lead_id,
            "ENTITY_TYPE": "lead",
            "COMMENT": f"🤖 ReplyFlow AI\n\n📝 Вопрос:\n{question}\n\n💡 Ответ:\n{answer}"
        }
    }
    
    import ssl
    ssl_context = ssl.create_default_context()
    ssl_context.check_hostname = False
    ssl_context.verify_mode = ssl.CERT_NONE
    connector = aiohttp.TCPConnector(ssl=ssl_context)
    
    try:
        async with aiohttp.ClientSession(connector=connector) as session:
            async with session.post(full_url, json=payload) as resp:
                result = await resp.json()
                print(f"[CRM] Bitrix24 response: {result}")
                return result.get("result", False)
    except Exception as e:
        print(f"[CRM] Bitrix24 error: {e}")
        return False


# ===== ВСПОМОГАТЕЛЬНЫЕ ФУНКЦИИ =====
async def get_current_user_optional(
    db: Session = Depends(get_db),
    creds: Optional[HTTPAuthorizationCredentials] = Depends(security)
) -> Optional[User]:
    if not creds:
        return None
    token = creds.credentials
    
    try:
        payload = decode_access_token(token)
        if payload and "sub" in payload:
            user_id = int(payload["sub"])
            user = db.query(User).filter(User.id == user_id).first()
            if user:
                return user
    except Exception:
        pass
    
    user = db.query(User).filter(User.api_key == token).first()
    if user:
        return user
    
    return None


@app.on_event("startup")
def startup_event():
    init_db()

@app.exception_handler(Exception)
async def exception_handler(request: Request, exc: Exception):
    error_msg = str(exc)
    log_error(f"Unhandled exception: {error_msg}", {"path": request.url.path})
    
    if "429" not in error_msg and "401" not in error_msg:
        send_telegram_message_sync(f"🚨 <b>Ошибка на сервере</b>\n\n📍 {request.url.path}\n❌ {error_msg[:200]}")
    
    return JSONResponse(
        status_code=400,
        content={"detail": error_msg}
    )


# ===== ОПРЕДЕЛЕНИЕ ЯЗЫКА =====
_detector = None

def get_detector():
    global _detector
    if _detector is None:
        _detector = LanguageDetectorBuilder.from_languages(Language.UKRAINIAN, Language.ENGLISH).build()
    return _detector

def detect_language(text: str) -> str:
    try:
        lang = detect(text)
        if lang == 'uk':
            return 'uk'
        else:
            return 'en'
    except Exception:
        return 'uk'

# ===== Groq клиент =====
groq_client = None
if GROQ_API_KEY:
    groq_client = AsyncGroq(api_key=GROQ_API_KEY)


async def detect_client_state(text: str, language: str = "uk") -> str:
    if not groq_client:
        return "interest"
    
    prompt = f"""
Analyze the customer's message and determine their state. Return ONLY one word.

Customer: "{text}"

Possible states: doubt, interest, objection, price_request, ready_to_buy
"""
    try:
        response = await groq_client.chat.completions.create(
            model="llama-3.1-8b-instant",
            messages=[{"role": "user", "content": prompt}],
            max_tokens=10,
            temperature=0.1
        )
        raw = response.choices[0].message.content.strip().lower()
        raw = re.sub(r'[^a-z_]', '', raw)
        
        valid_states = ["doubt", "interest", "objection", "price_request", "ready_to_buy"]
        for state in valid_states:
            if state in raw:
                return state
        return "interest"
    except Exception:
        return "interest"


def clean_response_text(text: str) -> str:
    if not text:
        return text
    text = re.sub(r'\*\*([^*]+)\*\*', r'\1', text)
    text = re.sub(r'^\*\*[^*]+\*\*:\s*', '', text, flags=re.MULTILINE)
    text = re.sub(r'^\d+\.\s*', '', text, flags=re.MULTILINE)
    text = re.sub(r'^[-*]\s*', '', text, flags=re.MULTILINE)
    return text.strip()


def format_search_results(results: List[Dict[str, Any]]) -> str:
    if not results:
        return ""
    
    text = "\n\n🔍 **Результати пошуку:**\n"
    for i, r in enumerate(results, 1):
        text += f"\n{i}. **{r['title']}**\n   📎 {r['url']}\n   📄 {r['content']}\n"
    
    return text


def add_links_to_response(response: str, results: List[Dict[str, Any]]) -> str:
    if not results:
        return response
    
    has_any_link = False
    for r in results:
        if r['url'] in response:
            has_any_link = True
            break
    
    if has_any_link:
        return response
    
    links_text = "\n\n📌 **Ось що вдалося знайти:**\n"
    for i, r in enumerate(results[:3], 1):
        title = r['title'][:80] + "..." if len(r['title']) > 80 else r['title']
        links_text += f"\n{i}. **{title}**\n   🔗 {r['url']}\n"
    
    links_text += "\n💡 *Якщо потрібно щось конкретніше — уточніть, будь ласка!*"
    
    return response + links_text


def parse_json(text, language: str = "uk"):
    default_messages = {
        "uk": "Відповідь не згенерована. Спробуйте ще",
        "en": "Response not generated. Try again"
    }
    default_msg = default_messages.get(language, default_messages["uk"])
    
    if not text:
        return [default_msg, default_msg, default_msg]
    
    def extract_string(obj, depth=0):
        if depth > 10:
            return default_msg
            
        if isinstance(obj, str):
            return obj if len(obj) > 0 else default_msg
        elif isinstance(obj, dict):
            for key in ['text', 'content', 'message', 'response', 'answer', '1', '2', '3', 'soft', 'selling', 'short']:
                if key in obj:
                    return extract_string(obj[key], depth + 1)
            for value in obj.values():
                result = extract_string(value, depth + 1)
                if result != default_msg:
                    return result
            return default_msg
        elif isinstance(obj, list):
            for item in obj:
                result = extract_string(item, depth + 1)
                if result != default_msg:
                    return result
            return default_msg
        else:
            return str(obj) if obj else default_msg
    
    clean_text = re.sub(r'```json\s*', '', text)
    clean_text = re.sub(r'```\s*', '', clean_text)
    
    start = clean_text.find("{")
    if start == -1:
        start = clean_text.find("[")
    
    if start != -1:
        clean_text = clean_text[start:]
    
    brace_count = 0
    in_string = False
    escape_next = False
    json_end = -1
    
    for i, char in enumerate(clean_text):
        if escape_next:
            escape_next = False
            continue
        if char == '\\':
            escape_next = True
            continue
        if char == '"' and not escape_next:
            in_string = not in_string
            continue
        if not in_string:
            if char == '{':
                brace_count += 1
            elif char == '}':
                brace_count -= 1
                if brace_count == 0:
                    json_end = i + 1
                    break
    
    if json_end != -1:
        json_str = clean_text[:json_end]
        
        json_str = re.sub(r',\s*}', '}', json_str)
        json_str = re.sub(r',\s*]', ']', json_str)
        json_str = re.sub(r'{\s*,', '{', json_str)
        json_str = re.sub(r'\[\s*,', '[', json_str)
        
        def fix_inner_quotes(match):
            content = match.group(1)
            content = content.replace('"', '\\"')
            content = content.replace('\n', '\\n')
            return f'"{content}"'
        
        json_str = re.sub(r'"([^"\\]*(\\.[^"\\]*)*)"', fix_inner_quotes, json_str)
        
        try:
            data = json.loads(json_str)
            
            if isinstance(data, list) and len(data) > 0:
                data = data[0]
            
            if "1" in data and "2" in data and "3" in data:
                return [
                    extract_string(data.get("1", default_msg)),
                    extract_string(data.get("2", default_msg)),
                    extract_string(data.get("3", default_msg)),
                ]
            
            if "soft" in data and "selling" in data and "short" in data:
                return [
                    extract_string(data.get("soft", default_msg)),
                    extract_string(data.get("selling", default_msg)),
                    extract_string(data.get("short", default_msg)),
                ]
            
            keys = list(data.keys())
            if len(keys) >= 3:
                return [
                    extract_string(data.get(keys[0], default_msg)),
                    extract_string(data.get(keys[1], default_msg)),
                    extract_string(data.get(keys[2], default_msg)),
                ]
            elif len(keys) == 2:
                val = extract_string(data.get(keys[0], default_msg))
                return [val, val, val]
            elif len(keys) == 1:
                val = extract_string(data.get(keys[0], default_msg))
                return [val, val, val]
                
        except json.JSONDecodeError as e:
            print(f"[PARSE] JSON decode error: {e}")
            pass
    
    lines = clean_text.strip().split('\n')
    response_texts = []
    for line in lines:
        match = re.match(r'^\s*\d+[\.\)]\s*(.+)$', line.strip())
        if match:
            response_texts.append(match.group(1).strip())
        elif line.strip().startswith('"') and line.strip().endswith('"'):
            try:
                val = json.loads(line.strip())
                if isinstance(val, str):
                    response_texts.append(val)
            except:
                pass
    
    if len(response_texts) >= 3:
        return response_texts[:3]
    elif len(response_texts) == 2:
        return [response_texts[0], response_texts[1], response_texts[1]]
    elif len(response_texts) == 1:
        return [response_texts[0], response_texts[0], response_texts[0]]
    
    if clean_text.strip():
        clean = re.sub(r'\{.*\}', '', clean_text, flags=re.DOTALL)
        clean = re.sub(r'\[.*\]', '', clean, flags=re.DOTALL)
        clean = clean.strip()
        if clean:
            return [clean[:200], clean[:200], clean[:200]]
    
    return [default_msg, default_msg, default_msg]


# ===== Демо =====
DEMO_LIMIT = config.DEMO_LIMIT
demo_ips = {}

def get_client_ip(request: Request) -> str:
    forwarded = request.headers.get("X-Forwarded-For")
    if forwarded:
        return forwarded.split(",")[0].strip()
    return request.client.host


def get_daily_limit(subscription_plan: str) -> int:
    limits = {
        "free": 10,
        "starter": 50,
        "professional": 150,
        "business": float('inf')
    }
    return limits.get(subscription_plan, 10)


# ===== AI МОДЕЛИ =====
from openai import OpenAI
OPENROUTER_API_KEY = config.OPENROUTER_API_KEY
HF_TOKEN = config.HF_TOKEN

client_openrouter = OpenAI(base_url=config.OPENROUTER_BASE_URL, api_key=OPENROUTER_API_KEY) if OPENROUTER_API_KEY else None

HF_URL = config.HF_URL
HF_HEADERS = {"Authorization": f"Bearer {HF_TOKEN}", "Content-Type": "application/json"} if HF_TOKEN else None

MODELS = []

if GROQ_API_KEY and groq_client:
    for model_name in config.GROQ_MODELS:
        MODELS.append({"type": "groq", "name": model_name, "client": groq_client})

if OPENROUTER_API_KEY and client_openrouter:
    for model_name in config.OPENROUTER_MODELS:
        MODELS.append({"type": "openrouter", "name": model_name, "client": client_openrouter})
        
async def call_model(model: Dict, prompt: str, max_tokens: int, temperature: float) -> Optional[str]:
    """Вызвать конкретную модель"""
    try:
        if model["type"] == "groq":
            response = await model["client"].chat.completions.create(
                model=model["name"],
                messages=[{"role": "user", "content": prompt}],
                max_tokens=max_tokens,
                temperature=temperature
            )
            text = response.choices[0].message.content
            return text if text and len(text) > 20 else None
                
        elif model["type"] == "openrouter":
            response = await asyncio.to_thread(
                model["client"].chat.completions.create,
                model=model["name"],
                messages=[{"role": "user", "content": prompt}],
                max_tokens=max_tokens,
                temperature=temperature
            )
            text = response.choices[0].message.content
            return text if text and len(text) > 20 else None
    except Exception as e:
        print(f"[ERROR] Model {model['name']} failed: {e}")
        return None
    
    return None


async def generate_with_fallback(prompt, subscription_plan="free", language="uk", max_tokens=500, temperature=0.7):
    """Генерация с балансировкой моделей, возвращает (текст, имя_использованной_модели)"""
    
    selected_model = select_model(MODELS, subscription_plan, language)
    
    if not selected_model:
        for model in MODELS:
            try:
                start_time = time.time()
                result = await call_model(model, prompt, max_tokens, temperature)
                response_time = time.time() - start_time
                if result:
                    record_model_result(model["name"], True, response_time)
                    return result, model["name"]
                else:
                    record_model_result(model["name"], False, response_time)
            except Exception:
                continue
        return "", None
    
    try:
        start_time = time.time()
        result = await call_model(selected_model, prompt, max_tokens, temperature)
        response_time = time.time() - start_time
        if result:
            record_model_result(selected_model["name"], True, response_time)
            return result, selected_model["name"]
        else:
            record_model_result(selected_model["name"], False, response_time)
    except Exception:
        pass
    
    for model in MODELS:
        if model["name"] == selected_model["name"]:
            continue
        try:
            start_time = time.time()
            result = await call_model(model, prompt, max_tokens, temperature)
            response_time = time.time() - start_time
            if result:
                record_model_result(model["name"], True, response_time)
                return result, model["name"]
        except Exception:
            continue
    
    return "", None


# ===== Pydantic Models =====
class Question(BaseModel):
    text: str
    objection: Optional[str] = None
    session_id: Optional[str] = None
    reset: Optional[bool] = False
    selected_script: Optional[str] = None


class ImproveRequest(BaseModel):
    text: str
    note: Optional[str] = ""


class EmailMessage(BaseModel):
    message: str


class CRMConnectRequest(BaseModel):
    crm_type: str
    api_key: str


class WebhookRegisterRequest(BaseModel):
    url: str
    events: List[str] = ["generation.created", "limit.reached"]


class PaymentWebhookData(BaseModel):
    user_id: int
    plan: str
    days: int = 30
    payment_id: Optional[str] = None


# ===== БАЗОВЫЕ ЭНДПОИНТЫ =====
@app.get("/")
async def root():
    return {"message": "OK"}


@app.get("/demo-remaining")
async def get_demo_remaining(request: Request):
    client_ip = get_client_ip(request)
    demo_count = demo_ips.get(client_ip, 0)
    remaining = max(0, DEMO_LIMIT - demo_count)
    return {"remaining": remaining}


@app.post("/demo-generate")
async def demo_generate(
    q: Question,
    request: Request,
    db: Session = Depends(get_db)
):
    language = detect_language(q.text)
    current_user = await get_current_user_optional(db, await security(request))
    
    if current_user:
        return await generate(q, db, current_user)
    
    client_ip = get_client_ip(request)
    demo_count = demo_ips.get(client_ip, 0)
    
    if demo_count >= DEMO_LIMIT:
        error_msg = {
            "uk": f"Ви використали всі {DEMO_LIMIT} безкоштовні демо-генерації. Зареєструйтесь для продовження.",
            "en": f"You have used all {DEMO_LIMIT} free demo generations. Register to continue."
        }.get(language, f"Ви використали всі {DEMO_LIMIT} безкоштовні демо-генерації.")
        return JSONResponse(status_code=429, content={"detail": error_msg})
    
    context = q.objection or ("не вказано" if language == "uk" else "not specified")
    
    prompt = build_generate_prompt(
        lang=language,
        role="free",
        history="",
        text=q.text,
        context=context,
        state="",
        search_results="",
        search_instruction="",
        knowledge_text=""
    )
    
    raw, used_model = await generate_with_fallback(prompt, subscription_plan="free", max_tokens=200, temperature=0.5)
    parts = parse_json(raw, language)
    
    demo_ips[client_ip] = demo_count + 1
    remaining = DEMO_LIMIT - (demo_count + 1)
    
    return {
        "best_response": parts[0],
        "other_responses": parts[1:3],
        "demo_remaining": remaining
    }


# ===== ОСНОВНОЙ ЭНДПОИНТ ГЕНЕРАЦИИ =====
@app.post("/generate")
async def generate(
    q: Question,
    db: Session = Depends(get_db),
    current_user: Optional[User] = Depends(get_current_user_optional)
):
    start_time = time.time()
    
    print(f"\n[REQUEST] User: {current_user.email if current_user else 'anonymous'}")
    print(f"[REQUEST] Text: {q.text[:100]}")
    
    language = detect_language(q.text)
    
    if not current_user:
        error_msg = {
            "uk": "Необхідна авторизація",
            "en": "Authorization required"
        }.get(language, "Необхідна авторизація")
        raise HTTPException(status_code=401, detail=error_msg)
    
    # ===== ПРОВЕРКА ПОДПИСКИ =====
    now = datetime.utcnow()
    
    if current_user.subscription_plan != "free" and current_user.subscription_end_date:
        if now > current_user.subscription_end_date:
            old_plan = current_user.subscription_plan
            current_user.subscription_plan = "free"
            current_user.subscription_end_date = None
            db.commit()
            print(f"[SUBSCRIPTION] User {current_user.email} downgraded from {old_plan} to free (expired)")
            
            error_msg = {
                "uk": f"Ваша підписка {old_plan} закінчилася. Будь ласка, оновіть тариф для продовження.",
                "en": f"Your {old_plan} subscription has expired. Please renew to continue."
            }.get(language, f"Your {old_plan} subscription has expired.")
            return JSONResponse(status_code=403, content={"detail": error_msg})
    
    # ===== ДНЕВНОЙ ЛИМИТ =====
    today_date = date.today()
    
    if current_user.last_generation_date is None:
        current_user.generations_today = 0
        current_user.last_generation_date = datetime.utcnow()
        current_user._daily_warning_sent = False
        db.commit()
        print(f"[RESET] First generation for user {current_user.email}")
    elif current_user.last_generation_date.date() != today_date:
        old_count = current_user.generations_today
        current_user.generations_today = 0
        current_user.last_generation_date = datetime.utcnow()
        current_user._daily_warning_sent = False
        db.commit()
        print(f"[RESET] Daily limit reset for user {current_user.email} from {old_count} to 0")
    
    daily_limit = get_daily_limit(current_user.subscription_plan)
    if current_user.generations_today >= daily_limit:
        error_msg = {
            "uk": f"Перевищено ліміт генерацій на сьогодні ({daily_limit}).",
            "en": f"You have exceeded the daily generation limit ({daily_limit})."
        }.get(language, f"Перевищено ліміт генерацій на сьогодні ({daily_limit}).")
        return JSONResponse(status_code=429, content={"detail": error_msg})
    
    # ===== ОПРЕДЕЛЯЕМ ПЛАН =====
    plan = current_user.subscription_plan
    
    # ===== ЖЕСТКИЕ ОГРАНИЧЕНИЯ ДЛЯ BUSINESS =====
    if plan == "business":
        BUSINESS_STRICT_DAILY_LIMIT = 5000
        if current_user.generations_today >= BUSINESS_STRICT_DAILY_LIMIT:
            error_msg = {
                "uk": f"Перевищено жорсткий денний ліміт для Business тарифу ({BUSINESS_STRICT_DAILY_LIMIT} генерацій). Зв'яжіться з підтримкою.",
                "en": f"Business daily limit exceeded ({BUSINESS_STRICT_DAILY_LIMIT} generations). Contact support."
            }.get(language, f"Business daily limit exceeded ({BUSINESS_STRICT_DAILY_LIMIT} generations).")
            return JSONResponse(status_code=429, content={"detail": error_msg})
        
        BUSINESS_STRICT_MONTHLY_LIMIT = 50000
        if current_user.total_generations >= BUSINESS_STRICT_MONTHLY_LIMIT:
            error_msg = {
                "uk": f"Перевищено жорсткий місячний ліміт для Business тарифу ({BUSINESS_STRICT_MONTHLY_LIMIT} генерацій). Зв'яжіться з підтримкою.",
                "en": f"Business monthly limit exceeded ({BUSINESS_STRICT_MONTHLY_LIMIT} generations). Contact support."
            }.get(language, f"Business monthly limit exceeded ({BUSINESS_STRICT_MONTHLY_LIMIT} generations).")
            return JSONResponse(status_code=429, content={"detail": error_msg})
        
        BUSINESS_STRICT_BUDGET = 200.0
        if current_user.total_cost_month >= BUSINESS_STRICT_BUDGET:
            error_msg = {
                "uk": f"Перевищено жорсткий бюджет для Business тарифу (${BUSINESS_STRICT_BUDGET}). Зв'яжіться з підтримкою.",
                "en": f"Business budget exceeded (${BUSINESS_STRICT_BUDGET}). Contact support."
            }.get(language, f"Business budget exceeded (${BUSINESS_STRICT_BUDGET}).")
            return JSONResponse(status_code=429, content={"detail": error_msg})
    
    session_id = q.session_id or str(uuid.uuid4())
    if q.reset:
        reset_memory(session_id)
    
    history_text = build_history_text(session_id, language)
    base_context = q.objection or ("не вказано" if language == "uk" else "not specified")
    
    # ===== ПОЛУЧАЕМ ВЫБРАННЫЙ СКРИПТ =====
    selected_script_id = q.selected_script
    script_template = None
    if selected_script_id:
        from app.models import SalesScript
        script = db.query(SalesScript).filter(SalesScript.id == selected_script_id).first()
        if script:
            script_template = script.template
            print(f"[SCRIPT] Using script: {script.name}")
    
    # ===== ПОЛУЧАЕМ ВСЕ СКРИПТЫ ПОЛЬЗОВАТЕЛЯ =====
    user_scripts = ""
    if plan in ["professional", "business"] and not script_template:
        from app.models import SalesScript
        scripts_list = db.query(SalesScript).filter(
            SalesScript.user_id == current_user.id,
            SalesScript.is_active == True
        ).all()
        
        if scripts_list:
            user_scripts = "\n\n📝 Доступные скрипты продаж:\n"
            for s in scripts_list:
                user_scripts += f"\n--- {s.name} ({s.category}) ---\n{s.template}\n"
    
    # ===== БАЗА ЗНАНИЙ =====
    knowledge_text = ""
    if plan in ["professional", "business"]:
        kb_items = db.query(KnowledgeBase).filter(
            KnowledgeBase.user_id == current_user.id,
            KnowledgeBase.is_active == True
        ).all()
        
        if kb_items:
            knowledge_parts = []
            for item in kb_items:
                knowledge_parts.append(f"=== {item.name} ===\n{item.content[:1000]}")
            knowledge_text = "\n\n📚 БАЗА ЗНАНИЙ (используй ЭТИ данные!):\n" + "\n\n".join(knowledge_parts)
            print(f"[KNOWLEDGE] Loaded {len(kb_items)} items for user {current_user.email}")
    
    # ===== ПРОВЕРКА БЮДЖЕТА =====
    estimated_cost = calculate_cost("llama-3.1-8b-instant")
    allowed, remaining_budget, monthly_limit = check_budget(current_user, estimated_cost)
    
    if not allowed:
        error_msg = {
            "uk": f"Перевищено місячний бюджет (${monthly_limit:.2f}). Будь ласка, оновіть тариф або зачекайте наступного місяця.",
            "en": f"Monthly budget exceeded (${monthly_limit:.2f}). Please upgrade your plan or wait for next month."
        }.get(language, "Monthly budget exceeded.")
        return JSONResponse(status_code=429, content={"detail": error_msg})
    
    remaining_percent = get_remaining_budget_percent(current_user)
    budget_warning = None
    if remaining_percent < 20 and plan != "business":
        budget_warning = {
            "uk": f"⚠️ Ви використали ${current_user.total_cost_month:.2f} з ${monthly_limit:.2f} місячного ліміту",
            "en": f"⚠️ You have used ${current_user.total_cost_month:.2f} of your ${monthly_limit:.2f} monthly limit"
        }.get(language, f"Used ${current_user.total_cost_month:.2f} of ${monthly_limit:.2f} monthly limit")
        
    # ===== ПРОВЕРКА КЭША =====
    cached_responses = get_cached_response(q.text, plan, language)
    if cached_responses:
        print(f"[CACHE] Hit for text: {q.text[:50]}...")
        parts = cached_responses
        
        current_user.total_generations += 1
        current_user.generations_today += 1
        db.commit()
        
        add_to_memory(session_id, q.text, parts[0])
        
        response_data = {
            "session_id": session_id,
            "best_response": parts[0],
            "other_responses": parts[1:3],
            "cached": True
        }
        
        if budget_warning:
            response_data["budget_warning"] = budget_warning
        
        # Аналитика для кэшированного ответа
        response_time_ms = int((time.time() - start_time) * 1000)
        tokens_used = len(parts[0]) // 4 if parts and parts[0] else 0
        
        try:
            from app.models import Analytics
            analytics = Analytics(
                user_id=current_user.id,
                user_email=current_user.email,
                subscription_plan=plan,
                question=q.text[:1000] if q.text else None,
                answer=parts[0][:1000] if parts and parts[0] else None,
                model_used="cached",
                response_time_ms=response_time_ms,
                tokens_used=tokens_used
            )
            db.add(analytics)
            db.commit()
        except Exception as e:
            print(f"[ANALYTICS] Error: {e}")
        
        print(f"[REQUEST] Completed (from cache)\n")
        return response_data
    
    # ===== АНАЛИЗ СОСТОЯНИЯ КЛИЕНТА =====
    client_state = None
    state_context = ""
    
    if plan in ["professional", "business"] and not script_template:
        client_state = await detect_client_state(q.text, language)
        client_state = re.sub(r'[^a-z_]', '', client_state.lower())
        valid_states = ["doubt", "interest", "objection", "price_request", "ready_to_buy"]
        if client_state not in valid_states:
            client_state = "interest"
        
        state_messages = {
            "uk": {
                "doubt": "Стан: сумнів (doubt)",
                "interest": "Стан: інтерес (interest)",
                "objection": "Стан: заперечення (objection)",
                "price_request": "Стан: запит ціни (price_request)",
                "ready_to_buy": "Стан: готовий купити (ready_to_buy)"
            },
            "en": {
                "doubt": "State: doubt",
                "interest": "State: interest",
                "objection": "State: objection",
                "price_request": "State: price request",
                "ready_to_buy": "State: ready to buy"
            }
        }
        state_context = state_messages[language].get(client_state, "")
    elif script_template:
        state_context = f"🚨 Используй этот скрипт: {script_template}"
    
    # ===== ПОШУК =====
    search_results = None
    search_results_text = ""
    search_instruction = ""
    
    if plan in ["professional", "business"] and should_search_external(q.text, language):
        print(f"[SEARCH] Searching...")
        search_results = await search_legislation(q.text, language)
        if search_results:
            search_results_text = format_search_results(search_results)
            search_instruction = "Обов'язково використовуй надані посилання у відповіді. Додавай їх як активні гіперпосилання в форматі [Назва](URL)." if language == "uk" else "Be sure to use the provided links in your response. Add them as active hyperlinks in the format [Title](URL)."
            print(f"[SEARCH] Found {len(search_results)} results")
        else:
            print(f"[SEARCH] No results found")
            search_results_text = "\n\n🔍 На жаль, нічого не знайдено. Спробуйте уточнити запит.\n"
            search_instruction = ""
    
    # ===== СКРИПТЫ (дополнительная инструкция) =====
    script_context = ""
    if selected_script_id and selected_script_id != "auto":
        from app.models import SalesScript
        script = db.query(SalesScript).filter(SalesScript.id == int(selected_script_id)).first()
        if script:
            script_context = f"\n\n🚨 Используй ЭТОТ скрипт как основу для ответа:\n{script.template}\nАдаптируй его под конкретную ситуацию клиента.\n"
            print(f"[SCRIPT] Using script: {script.name}")
    elif plan in ["professional", "business"]:
        from app.models import SalesScript
        scripts_list = db.query(SalesScript).filter(
            SalesScript.user_id == current_user.id,
            SalesScript.is_active == True
        ).all()
        if scripts_list:
            user_scripts_text = "\n\n📝 Доступные скрипты продаж:\n"
            for s in scripts_list:
                user_scripts_text += f"\n--- {s.name} ({s.category}) ---\n{s.template}\n"
            script_context = user_scripts_text + "\n🚨 Выбери скрипт, который лучше всего подходит к ситуации клиента. Адаптируй его под конкретный запрос.\n"
    
    # ===== ОПРЕДЕЛЯЕМ РОЛЬ ДЛЯ ПРОМПТА =====
    if plan == "business":
        role = "business"
    elif plan == "professional":
        role = "professional"
    else:
        role = "free"
    
    # ===== ФОРМИРУЕМ ПРОМПТ =====
    prompt = build_generate_prompt(
        lang=language,
        role=role,
        history=history_text,
        text=q.text,
        context=base_context,
        state=state_context,
        search_results=search_results_text,
        search_instruction=search_instruction + script_context,
        knowledge_text=knowledge_text
    )
    
    # ===== ПАРАМЕТРЫ ГЕНЕРАЦИИ =====
    if plan == "business":
        temperature = 0.3
        base_tokens = 600
    elif plan == "professional":
        temperature = 0.5
        base_tokens = 500
    else:
        temperature = 0.5
        base_tokens = 400
    
    # Генерація
    default_msg = "Дякуємо за ваше звернення! Наш менеджер зв'яжеться з вами найближчим часом." if language == "uk" else "Thank you for your inquiry! Our manager will contact you shortly."
    
    max_attempts = 3
    raw = None
    parts = None
    used_model = None
    
    for attempt in range(max_attempts):
        current_tokens = base_tokens + (attempt * 200)
        
        raw, used_model = await generate_with_fallback(prompt, subscription_plan=plan, language=language, max_tokens=current_tokens, temperature=temperature)
        
        if not raw or len(raw) < 20:
            continue
        
        parts = parse_json(raw, language)
        
        if parts and parts[0] != default_msg and len(parts[0]) > 20:
            break
    
    if not used_model:
        used_model = "llama-3.1-8b-instant"
    
    if not parts or not parts[0] or parts[0] == default_msg:
        if raw and len(raw) > 20:
            parts = parse_json(raw, language)
        else:
            parts = [default_msg, default_msg, default_msg]
    
    # Додаємо посилання
    if search_results:
        parts[0] = add_links_to_response(parts[0], search_results)
        parts[1] = add_links_to_response(parts[1], search_results)
        parts[2] = add_links_to_response(parts[2], search_results)
    
    # Сохраняем в кэш
    set_cached_response(q.text, plan, language, parts)
    
    add_to_memory(session_id, q.text, parts[0])
    
    # Обновляем счётчики
    current_user.total_generations += 1
    current_user.generations_today += 1
    
    # Сброс месячных флагов при смене месяца
    if current_user.cost_reset_date is None:
        current_user.cost_reset_date = datetime.utcnow()
        current_user.total_cost_month = 0
        current_user._budget_warning_sent = False
        current_user._monthly_warning_sent = False
    elif datetime.utcnow().month != current_user.cost_reset_date.month or datetime.utcnow().year != current_user.cost_reset_date.year:
        current_user.total_cost_month = 0
        current_user.cost_reset_date = datetime.utcnow()
        current_user._budget_warning_sent = False
        current_user._monthly_warning_sent = False
        print(f"[RESET] Monthly budget reset for user {current_user.email}")
    
    # Обновляем стоимость с использованием реальной модели
    actual_cost = calculate_cost(used_model, input_tokens=500, output_tokens=len(parts[0]) if parts else 500)
    update_user_cost(current_user, actual_cost, db)
    
    db.commit()
    
    # ===== ОТПРАВКА В CRM =====
    if current_user.crm_type == "bitrix24" and current_user.crm_webhook:
        lead_id = q.session_id or str(current_user.id)
        await send_to_bitrix24(
            current_user.crm_webhook,
            lead_id,
            q.text,
            parts[0]
        )
    
    # ===== ОТПРАВКА В WEBHOOK =====
    if current_user.webhook_url:
        try:
            webhook_payload = {
                "event": "generation.created",
                "user_id": current_user.id,
                "user_email": current_user.email,
                "question": q.text,
                "answer": parts[0],
                "timestamp": datetime.utcnow().isoformat()
            }
            async with aiohttp.ClientSession() as session:
                await session.post(current_user.webhook_url, json=webhook_payload, timeout=3.0)
        except Exception as e:
            print(f"[WEBHOOK] Error: {e}")
    
    # ===== АНАЛИТИКА =====
    response_time_ms = int((time.time() - start_time) * 1000)
    tokens_used = len(parts[0]) // 4 if parts and parts[0] else 0
    
    try:
        from app.models import Analytics
        analytics = Analytics(
            user_id=current_user.id,
            user_email=current_user.email,
            subscription_plan=plan,
            question=q.text[:1000] if q.text else None,
            answer=parts[0][:1000] if parts and parts[0] else None,
            model_used=used_model,
            response_time_ms=response_time_ms,
            tokens_used=tokens_used
        )
        db.add(analytics)
        db.commit()
        print(f"[ANALYTICS] Saved: user={current_user.email}, model={used_model}, time={response_time_ms}ms")
    except Exception as e:
        print(f"[ANALYTICS] Error: {e}")
    
    response_data = {
        "session_id": session_id,
        "best_response": parts[0],
        "other_responses": parts[1:3]
    }
    
    if budget_warning:
        response_data["budget_warning"] = budget_warning
    
    if plan in ["professional", "business"] and client_state:
        response_data["client_state"] = client_state
    
    print(f"[REQUEST] Completed with model: {used_model}, time: {response_time_ms}ms\n")
    return response_data
    
# ===== СТРИМИНГ ГЕНЕРАЦИИ =====
@app.post("/generate-stream")
async def generate_stream(
    q: Question,
    db: Session = Depends(get_db),
    current_user: Optional[User] = Depends(get_current_user_optional)
):
    """Генерация ответа с потоковой передачей (только первый ответ)"""
    
    print(f"\n[STREAM] User: {current_user.email if current_user else 'anonymous'}")
    print(f"[STREAM] Text: {q.text[:100]}")
    
    language = detect_language(q.text)
    
    if not current_user:
        error_msg = {
            "uk": "Необхідна авторизація",
            "en": "Authorization required"
        }.get(language, "Необхідна авторизація")
        raise HTTPException(status_code=401, detail=error_msg)
    
    # Проверка дневного лимита
    daily_limit = get_daily_limit(current_user.subscription_plan)
    if current_user.generations_today >= daily_limit:
        error_msg = {
            "uk": f"Перевищено ліміт генерацій на сьогодні ({daily_limit}).",
            "en": f"You have exceeded the daily generation limit ({daily_limit})."
        }.get(language, f"Перевищено ліміт генерацій на сьогодні ({daily_limit}).")
        raise HTTPException(status_code=429, detail=error_msg)
    
    session_id = q.session_id or str(uuid.uuid4())
    if q.reset:
        reset_memory(session_id)
    
    history_text = build_history_text(session_id, language)
    base_context = q.objection or ("не вказано" if language == "uk" else "not specified")
    
    plan = current_user.subscription_plan
    
    # База знаний
    knowledge_text = ""
    if plan in ["professional", "business"]:
        kb_items = db.query(KnowledgeBase).filter(
            KnowledgeBase.user_id == current_user.id,
            KnowledgeBase.is_active == True
        ).all()
        
        if kb_items:
            knowledge_parts = []
            for item in kb_items:
                knowledge_parts.append(f"=== {item.name} ===\n{item.content[:1000]}")
            knowledge_text = "\n\n📚 БАЗА ЗНАНИЙ (используй ЭТИ данные!):\n" + "\n\n".join(knowledge_parts)
    
    # Определяем роль
    if plan == "business":
        role = "business"
    elif plan == "professional":
        role = "professional"
    else:
        role = "free"
    
    # Анализ состояния клиента
    client_state = None
    state_context = ""
    if plan in ["professional", "business"]:
        client_state = await detect_client_state(q.text, language)
        client_state = re.sub(r'[^a-z_]', '', client_state.lower())
        valid_states = ["doubt", "interest", "objection", "price_request", "ready_to_buy"]
        if client_state not in valid_states:
            client_state = "interest"
        
        state_messages = {
            "uk": {
                "doubt": "Стан: сумнів (doubt)",
                "interest": "Стан: інтерес (interest)",
                "objection": "Стан: заперечення (objection)",
                "price_request": "Стан: запит ціни (price_request)",
                "ready_to_buy": "Стан: готовий купити (ready_to_buy)"
            },
            "en": {
                "doubt": "State: doubt",
                "interest": "State: interest",
                "objection": "State: objection",
                "price_request": "State: price request",
                "ready_to_buy": "State: ready to buy"
            }
        }
        state_context = state_messages[language].get(client_state, "")
    
    # Скрипты
    script_context = ""
    selected_script_id = q.selected_script
    
    if selected_script_id and selected_script_id != "auto":
        from app.models import SalesScript
        script = db.query(SalesScript).filter(SalesScript.id == int(selected_script_id)).first()
        if script:
            script_context = f"\n\n🚨 Используй ЭТОТ скрипт как основу для ответа:\n{script.template}\nАдаптируй его под конкретную ситуацию клиента.\n"
    elif plan in ["professional", "business"]:
        from app.models import SalesScript
        scripts_list = db.query(SalesScript).filter(
            SalesScript.user_id == current_user.id,
            SalesScript.is_active == True
        ).all()
        if scripts_list:
            user_scripts_text = "\n\n📝 Доступные скрипты продаж:\n"
            for s in scripts_list:
                user_scripts_text += f"\n--- {s.name} ({s.category}) ---\n{s.template}\n"
            script_context = user_scripts_text + "\n🚨 Выбери скрипт, который лучше всего подходит к ситуации клиента.\n"
    
    # Формируем промпт
    prompt = build_stream_prompt(
        lang=language,
        role=role,
        history=history_text,
        text=q.text,
        context=base_context,
        state=state_context,
        search_results="",
        search_instruction=script_context,
        knowledge_text=knowledge_text
    )
    
    temperature = 0.5
    
    async def stream_generator():
        full_response = ""
        try:
            if groq_client:
                stream = await groq_client.chat.completions.create(
                    model="llama-3.1-8b-instant",
                    messages=[{"role": "user", "content": prompt}],
                    max_tokens=500,
                    temperature=temperature,
                    stream=True
                )
                
                async for chunk in stream:
                    if chunk.choices[0].delta.content:
                        content = chunk.choices[0].delta.content
                        full_response += content
                        yield content
                
                add_to_memory(session_id, q.text, full_response)
                
                current_user.total_generations += 1
                current_user.generations_today += 1
                actual_cost = calculate_cost("llama-3.1-8b-instant", input_tokens=300, output_tokens=len(full_response))
                update_user_cost(current_user, actual_cost, db)
                db.commit()
                
            else:
                raw, used_model = await generate_with_fallback(prompt, subscription_plan=plan, language=language)
                add_to_memory(session_id, q.text, raw)
                
                current_user.total_generations += 1
                current_user.generations_today += 1
                if used_model:
                    actual_cost = calculate_cost(used_model, input_tokens=300, output_tokens=len(raw))
                    update_user_cost(current_user, actual_cost, db)
                db.commit()
                
                yield raw
                
        except Exception as e:
            print(f"[STREAM] Error: {e}")
            error_msg = "Помилка генерації. Спробуйте ще раз." if language == "uk" else "Generation error. Try again."
            yield error_msg
    
    return StreamingResponse(
        stream_generator(),
        media_type="text/plain",
        headers={
            "X-Session-Id": session_id,
            "Cache-Control": "no-cache",
            "Connection": "keep-alive",
        }
    )

# ===== АСИНХРОННАЯ ГЕНЕРАЦИЯ =====
@app.post("/generate-async")
async def generate_async(
    q: Question,
    db: Session = Depends(get_db),
    current_user: Optional[User] = Depends(get_current_user_optional)
):
    if not current_user:
        raise HTTPException(status_code=401, detail="Authorization required")
    
    task = generate_response_task.delay(
        prompt=q.text,
        subscription_plan=current_user.subscription_plan,
        language=detect_language(q.text),
        max_tokens=500,
        temperature=0.7
    )
    
    return {
        "task_id": task.id,
        "status": "processing",
        "message": "Task created. Poll /task-status/{task_id} for result."
    }


@app.get("/task-status/{task_id}")
async def get_task_status(task_id: str):
    result = get_task_result(task_id)
    
    if not result:
        return {"status": "not_found", "task_id": task_id}
    
    return {
        "task_id": task_id,
        "status": result.get("status"),
        "progress": result.get("progress", 0),
        "result": result.get("result") if result.get("status") == "completed" else None,
        "error": result.get("error") if result.get("status") == "failed" else None
    }

# ===== УЛУЧШЕНИЕ ОТВЕТА =====
@app.post("/improve-answer")
async def improve(
    req: ImproveRequest,
    db: Session = Depends(get_db),
    current_user: Optional[User] = Depends(get_current_user_optional)
):
    language = detect_language(req.text)
    
    if not current_user:
        error_msg = {
            "uk": "Необхідна авторизація",
            "en": "Authorization required"
        }.get(language, "Необхідна авторизація")
        raise HTTPException(status_code=401, detail=error_msg)
    
    daily_limit = get_daily_limit(current_user.subscription_plan)
    if current_user.generations_today >= daily_limit:
        error_msg = {
            "uk": f"Перевищено ліміт генерацій на сьогодні ({daily_limit}).",
            "en": f"You have exceeded the daily generation limit ({daily_limit})."
        }.get(language, f"Перевищено ліміт генерацій на сьогодні ({daily_limit}).")
        return JSONResponse(status_code=429, content={"detail": error_msg})
    
    plan = current_user.subscription_plan
    if plan == "business":
        max_tokens = 400
        temperature = 0.3
    elif plan == "professional":
        max_tokens = 300
        temperature = 0.5
    else:
        max_tokens = 200
        temperature = 0.5
    
    context = req.note or ("не вказано" if language == "uk" else "not specified")
    
    prompt = build_improve_prompt(
        text=req.text,
        context=context,
        lang=language
    )
    
    result, used_model = await generate_with_fallback(prompt, subscription_plan=plan, max_tokens=max_tokens, temperature=temperature)
    
    if not used_model:
        used_model = "llama-3.1-8b-instant"
    
    if result:
        result = result.strip()
        result = result.strip('"').strip("'")
        result = clean_response_text(result)
    
    # Обновляем счетчики
    today = date.today()
    if current_user.last_generation_date is None or current_user.last_generation_date.date() != today:
        current_user.generations_today = 0
        current_user.last_generation_date = datetime.utcnow()
    
    current_user.total_generations += 1
    current_user.generations_today += 1
    actual_cost = calculate_cost(used_model, input_tokens=200, output_tokens=len(result) if result else 200)
    update_user_cost(current_user, actual_cost, db)
    db.commit()
    
    default_msg = {
        "uk": "Не вдалося покращити відповідь. Спробуйте ще раз.",
        "en": "Failed to improve the answer. Try again."
    }.get(language, "Failed to improve the answer. Try again.")
    
    return {"improved_text": result or default_msg}


# ===== ОТПРАВКА EMAIL =====
@app.post("/send-email")
async def send_email(
    data: EmailMessage,
    request: Request,
    current_user: Optional[User] = Depends(get_current_user_optional)
):
    try:
        client_ip = get_client_ip(request)
        user_agent = request.headers.get("user-agent", "Unknown")
        
        email_content = f"""
📨 НОВОЕ СООБЩЕНИЕ С ЛЕНДИНГА

━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━

📝 ТЕКСТ СООБЩЕНИЯ:
{data.message}

━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━

👤 ИНФОРМАЦИЯ О ПОЛЬЗОВАТЕЛЕ:
"""
        
        if current_user:
            email_content += f"""
✅ Пользователь АВТОРИЗОВАН:
   • ID: {current_user.id}
   • Email: {current_user.email}
   • Username: {current_user.username}
   • Тариф: {current_user.subscription_plan}
"""
        else:
            email_content += f"""
⚠️ Пользователь НЕ АВТОРИЗОВАН (аноним)
"""
        
        email_content += f"""
━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━

🌐 ТЕХНИЧЕСКАЯ ИНФОРМАЦИЯ:
   • IP адрес: {client_ip}
   • User-Agent: {user_agent}
   • Время: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}

━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
"""
        
        # Пытаемся отправить через SMTP
        try:
            msg = MIMEText(email_content, "plain", "utf-8")
            msg["Subject"] = Header(f"Новый отзыв с ReplyFlow AI", "utf-8")
            msg["From"] = config.EMAIL_ADDRESS
            msg["To"] = config.EMAIL_ADDRESS
            
            server = smtplib.SMTP_SSL(config.EMAIL_SMTP_HOST, config.EMAIL_SMTP_PORT)
            server.login(config.EMAIL_ADDRESS, config.EMAIL_PASSWORD)
            server.send_message(msg)
            server.quit()
            
            return {"status": "success", "message": "Спасибо за отзыв!"}
        except Exception as smtp_error:
            # Если SMTP не работает — пробуем отправить через Mailgun API
            print(f"SMTP failed: {smtp_error}, trying Mailgun API...")
            
            api_key = os.environ.get("MAILGUN_API_KEY")
            if not api_key:
                raise Exception("Mailgun API key not configured")
            
            domain = "sandbox95ed721e2bf445488e3a9c910019da16.mailgun.org"
            
            import requests
            response = requests.post(
                f"https://api.mailgun.net/v3/{domain}/messages",
                auth=("api", api_key),
                data={
                    "from": f"ReplyFlow Bot <postmaster@{domain}>",
                    "to": [config.EMAIL_ADDRESS],  # отправляем на твой email
                    "subject": "Новое сообщение с сайта",
                    "text": email_content
                }
            )
            
            if response.status_code == 200:
                return {"status": "success", "message": "Спасибо за отзыв!"}
            else:
                raise Exception(f"Mailgun API error: {response.text}")
                
    except Exception as e:
        print(f"Ошибка отправки email: {e}")
        return {"status": "error", "message": str(e)}

# ===== ВЕБХУК ДЛЯ ОПЛАТЫ =====
@app.post("/api/webhooks/payment")
async def payment_webhook(
    data: PaymentWebhookData,
    db: Session = Depends(get_db)
):
    user = db.query(User).filter(User.id == data.user_id).first()
    if not user:
        raise HTTPException(status_code=404, detail="User not found")
    
    valid_plans = ["starter", "professional", "business"]
    if data.plan not in valid_plans:
        raise HTTPException(status_code=400, detail=f"Invalid plan")
    
    user.subscription_plan = data.plan
    user.subscription_end_date = datetime.utcnow() + timedelta(days=data.days)
    user.subscription_auto_renew = True
    user.generations_today = 0
    user.total_cost_month = 0
    user.last_generation_date = datetime.utcnow()
    user.cost_reset_date = datetime.utcnow()
    db.commit()
    
    return {
        "status": "ok",
        "message": f"Subscription activated: {data.plan} until {user.subscription_end_date.date()}",
        "user_id": user.id,
        "plan": user.subscription_plan,
        "end_date": user.subscription_end_date.isoformat()
    }


# ===== ТЕСТОВЫЙ ЭНДПОИНТ ДЛЯ ИМИТАЦИИ ОПЛАТЫ =====
@app.post("/api/test/simulate-payment")
async def simulate_payment(
    user_id: int,
    plan: str = "starter",
    days: int = 30,
    db: Session = Depends(get_db)
):
    user = db.query(User).filter(User.id == user_id).first()
    if not user:
        raise HTTPException(status_code=404, detail="User not found")
    
    user.subscription_plan = plan
    user.subscription_end_date = datetime.utcnow() + timedelta(days=days)
    user.subscription_auto_renew = True
    user.generations_today = 0
    user.total_cost_month = 0
    db.commit()
    
    return {
        "status": "ok",
        "message": f"TEST: Subscription activated for user {user.email}",
        "plan": plan,
        "end_date": user.subscription_end_date.isoformat()
    }


# ===== ЭНДПОИНТЫ ИНТЕГРАЦИЙ =====
@app.post("/api/integrations/crm/connect")
async def connect_crm(
    data: CRMConnectRequest,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    if current_user.subscription_plan != "business":
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="CRM integration available only for Business plan"
        )
    
    is_valid = False
    crm_data = {}
    
    if data.crm_type == "hubspot":
        is_valid = await verify_hubspot_api_key(data.api_key)
        if is_valid:
            crm_data = {"api_key": data.api_key}
    elif data.crm_type == "amocrm":
        parts = data.api_key.split(":")
        if len(parts) == 2:
            subdomain, api_key = parts
            is_valid = await verify_amocrm_api_key(api_key, subdomain)
            if is_valid:
                crm_data = {"subdomain": subdomain, "api_key": api_key}
        else:
            raise HTTPException(status_code=400, detail="Invalid amoCRM format. Use 'subdomain:api_key'")
    elif data.crm_type == "pipedrive":
        is_valid = await verify_pipedrive_api_key(data.api_key)
        if is_valid:
            crm_data = {"api_token": data.api_key}
    elif data.crm_type == "bitrix24":
        is_valid = True
        crm_data = {"webhook_url": data.api_key}
    else:
        raise HTTPException(status_code=400, detail=f"Unknown CRM type: {data.crm_type}")
    
    if not is_valid:
        raise HTTPException(status_code=400, detail=f"Invalid {data.crm_type} API key")
    
    current_user.crm_type = data.crm_type
    current_user.crm_api_key = data.api_key
    current_user.crm_data = crm_data
    db.commit()
    
    return {
        "status": "ok",
        "message": f"CRM {data.crm_type} connected successfully"
    }


@app.delete("/api/integrations/crm/disconnect")
async def disconnect_crm(
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    if current_user.subscription_plan != "business":
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="CRM integration available only for Business plan"
        )
    
    current_user.crm_type = None
    current_user.crm_api_key = None
    current_user.crm_webhook = None
    current_user.crm_data = None
    db.commit()
    
    return {"status": "ok", "message": "CRM disconnected"}


@app.post("/api/integrations/webhooks/register")
async def register_webhook(
    data: WebhookRegisterRequest,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    if current_user.subscription_plan != "business":
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Webhooks available only for Business plan"
        )
    
    # Проверяем URL тестовым уведомлением
    test_payload = {
        "event": "test",
        "message": "Webhook configured successfully",
        "user_id": current_user.id,
        "timestamp": datetime.utcnow().isoformat()
    }
    
    try:
        async with aiohttp.ClientSession() as session:
            async with session.post(data.url, json=test_payload, timeout=5.0) as resp:
                if resp.status not in [200, 201, 202, 204]:
                    raise HTTPException(status_code=400, detail="Webhook URL did not acknowledge test notification")
    except asyncio.TimeoutError:
        raise HTTPException(status_code=400, detail="Webhook URL timeout")
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Webhook test failed: {str(e)}")
    
    current_user.webhook_url = data.url
    db.commit()
    
    return {
        "status": "ok",
        "message": f"Webhook registered for {data.url}",
        "events": data.events
    }


@app.delete("/api/integrations/webhooks/unregister")
async def unregister_webhook(
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    if current_user.subscription_plan != "business":
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Webhooks available only for Business plan"
        )
    
    current_user.webhook_url = None
    db.commit()
    
    return {"status": "ok", "message": "Webhook unregistered"}


@app.get("/api/integrations/status")
async def get_integrations_status(
    current_user: User = Depends(get_current_user)
):
    if current_user.subscription_plan != "business":
        return {
            "plan": current_user.subscription_plan,
            "integrations": {
                "crm": {"available": False, "message": "Upgrade to Business plan"},
                "bitrix24": {"available": False, "message": "Upgrade to Business plan"},
                "webhooks": {"available": False, "message": "Upgrade to Business plan"},
                "api": {"available": False, "message": "Upgrade to Business plan"}
            }
        }
    
    return {
        "plan": current_user.subscription_plan,
        "integrations": {
            "crm": {
                "available": True,
                "connected": bool(current_user.crm_type and current_user.crm_type != "bitrix24"),
                "crm_type": current_user.crm_type,
                "message": "CRM integration"
            },
            "bitrix24": {
                "available": True,
                "connected": bool(current_user.crm_type == "bitrix24"),
                "message": "Bitrix24 integration"
            },
            "webhooks": {
                "available": True,
                "connected": bool(current_user.webhook_url),
                "webhook_url": current_user.webhook_url,
                "message": "Webhooks"
            },
            "api": {
                "available": True,
                "connected": bool(current_user.api_key),
                "message": "API access"
            }
        }
    }


# ===== API КЛЮЧИ =====
@app.post("/api/integrations/api-keys/generate")
async def generate_api_key(
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    if current_user.subscription_plan not in ["starter", "professional", "business"]:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="API access available for Starter, Professional and Business plans"
        )
    
    import secrets
    api_key = f"rf_{secrets.token_urlsafe(32)}"
    
    current_user.api_key = api_key
    db.commit()
    
    return {
        "status": "ok",
        "api_key": api_key,
        "message": "API key generated"
    }


@app.delete("/api/integrations/api-keys/revoke")
async def revoke_api_key(
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    if current_user.subscription_plan not in ["starter", "professional", "business"]:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="API access available for Starter, Professional and Business plans"
        )
    
    current_user.api_key = None
    db.commit()
    
    return {"status": "ok", "message": "API key revoked"}


# ===== ЗАГРУЗКА ФАЙЛОВ В БАЗУ ЗНАНИЙ =====
@app.post("/api/knowledge/upload-file")
async def upload_knowledge_file(
    file: UploadFile = File(...),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    from app.models import KnowledgeBase
    
    if current_user.subscription_plan not in ["professional", "business"]:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Knowledge base available for Professional and Business plans"
        )
    
    file_ext = file.filename.split(".")[-1].lower()
    if file_ext not in ["pdf", "docx", "xlsx", "txt"]:
        raise HTTPException(status_code=400, detail="Unsupported file type")
    
    safe_filename = f"{uuid.uuid4()}_{file.filename}"
    file_path = os.path.join(UPLOAD_DIR, safe_filename)
    
    async with aiofiles.open(file_path, "wb") as f:
        content = await file.read()
        await f.write(content)
    
    extracted_text = extract_text_from_file(file_path, file_ext)
    
    kb_item = KnowledgeBase(
        user_id=current_user.id,
        name=file.filename,
        content=extracted_text,
        file_type=file_ext,
        is_active=True
    )
    db.add(kb_item)
    db.commit()
    db.refresh(kb_item)
    
    return {
        "id": kb_item.id,
        "name": kb_item.name,
        "file_type": kb_item.file_type,
        "content_preview": extracted_text[:200],
        "message": "File uploaded"
    }


# ===== ROUTERS =====
app.include_router(users_router)
app.include_router(admin_router)
app.include_router(password_reset_router)
app.include_router(scripts_router)
app.include_router(knowledge_router)
app.include_router(bot_router)
