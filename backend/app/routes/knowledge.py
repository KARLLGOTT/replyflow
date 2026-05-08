from fastapi import APIRouter, Depends, HTTPException, status, UploadFile, File
from sqlalchemy.orm import Session
from typing import List, Optional
from pydantic import BaseModel
from datetime import datetime
import os
import uuid
import aiofiles
from app.database import get_db
from app.models import User, KnowledgeBase
from app.dependencies import get_current_user

router = APIRouter(prefix="/api/knowledge", tags=["Knowledge Base"])

UPLOAD_DIR = "uploads/knowledge"
os.makedirs(UPLOAD_DIR, exist_ok=True)


def extract_text_from_file(file_path: str, file_type: str) -> str:
    """Извлечь текст из файла"""
    import PyPDF2
    from docx import Document
    import openpyxl
    
    try:
        if file_type == "pdf":
            with open(file_path, "rb") as f:
                reader = PyPDF2.PdfReader(f)
                text = "\n".join([page.extract_text() for page in reader.pages if page.extract_text()])
            return text[:5000] if text else "PDF не содержит текста"
        
        elif file_type == "docx":
            doc = Document(file_path)
            text = "\n".join([para.text for para in doc.paragraphs if para.text])
            return text[:5000]
        
        elif file_type == "xlsx":
            wb = openpyxl.load_workbook(file_path, data_only=True)
            text = ""
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    text += " ".join([str(cell) for cell in row if cell]) + "\n"
            return text[:5000]
        
        elif file_type == "txt":
            with open(file_path, "r", encoding="utf-8") as f:
                return f.read()[:5000]
        
        else:
            return f"Неподдерживаемый тип: {file_type}"
    except Exception as e:
        return f"Ошибка извлечения: {e}"


# ===== SCHEMAS =====
class KnowledgeCreate(BaseModel):
    name: str
    content: str
    file_type: str = "text"


class KnowledgeUpdate(BaseModel):
    name: Optional[str] = None
    content: Optional[str] = None
    is_active: Optional[bool] = None


class KnowledgeResponse(BaseModel):
    id: int
    name: str
    content: str
    file_type: str
    is_active: bool
    created_at: datetime
    updated_at: datetime

    class Config:
        from_attributes = True


# ===== ЭНДПОИНТЫ =====
@router.get("/", response_model=List[KnowledgeResponse])
def get_knowledge(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    items = db.query(KnowledgeBase).filter(
        KnowledgeBase.user_id == current_user.id
    ).order_by(KnowledgeBase.created_at.desc()).all()
    return items


@router.post("/", response_model=KnowledgeResponse)
def create_knowledge(
    data: KnowledgeCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    if current_user.subscription_plan not in ["professional", "business"]:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Knowledge base available for Professional and Business plans"
        )
    
    if not data.content or len(data.content.strip()) < 10:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Content must be at least 10 characters"
        )
    
    kb_item = KnowledgeBase(
        user_id=current_user.id,
        name=data.name,
        content=data.content,
        file_type=data.file_type,
        is_active=True
    )
    db.add(kb_item)
    db.commit()
    db.refresh(kb_item)
    return kb_item


@router.post("/upload-file")
async def upload_knowledge_file(
    file: UploadFile = File(...),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """Загрузить файл (PDF, DOCX, XLSX, TXT) в базу знаний"""
    if current_user.subscription_plan not in ["professional", "business"]:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Knowledge base available for Professional and Business plans"
        )
    
    # Сохраняем файл
    file_ext = file.filename.split(".")[-1].lower()
    if file_ext not in ["pdf", "docx", "xlsx", "txt"]:
        raise HTTPException(status_code=400, detail="Unsupported file type")
    
    safe_filename = f"{uuid.uuid4()}_{file.filename}"
    file_path = os.path.join(UPLOAD_DIR, safe_filename)
    
    async with aiofiles.open(file_path, "wb") as f:
        content = await file.read()
        await f.write(content)
    
    # Извлекаем текст
    extracted_text = extract_text_from_file(file_path, file_ext)
    
    # Сохраняем в БД
    kb_item = KnowledgeBase(
        user_id=current_user.id,
        name=file.filename,
        content=extracted_text,
        file_type=file_ext,
        is_active=True
    )
    db.add(kb_item)
    db.commit()
    db.refresh(kb_item)
    
    return {
        "id": kb_item.id,
        "name": kb_item.name,
        "file_type": kb_item.file_type,
        "content_preview": extracted_text[:200],
        "message": "File uploaded"
    }


@router.put("/{item_id}", response_model=KnowledgeResponse)
def update_knowledge(
    item_id: int,
    data: KnowledgeUpdate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    item = db.query(KnowledgeBase).filter(
        KnowledgeBase.id == item_id,
        KnowledgeBase.user_id == current_user.id
    ).first()
    
    if not item:
        raise HTTPException(status_code=404, detail="Item not found")
    
    if data.name is not None:
        item.name = data.name
    if data.content is not None:
        item.content = data.content
    if data.is_active is not None:
        item.is_active = data.is_active
    
    item.updated_at = datetime.utcnow()
    db.commit()
    db.refresh(item)
    return item


@router.delete("/{item_id}")
def delete_knowledge(
    item_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    item = db.query(KnowledgeBase).filter(
        KnowledgeBase.id == item_id,
        KnowledgeBase.user_id == current_user.id
    ).first()
    
    if not item:
        raise HTTPException(status_code=404, detail="Item not found")
    
    db.delete(item)
    db.commit()
    
    return {"status": "ok", "message": "Item deleted"}


@router.post("/{item_id}/toggle")
def toggle_knowledge(
    item_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    item = db.query(KnowledgeBase).filter(
        KnowledgeBase.id == item_id,
        KnowledgeBase.user_id == current_user.id
    ).first()
    
    if not item:
        raise HTTPException(status_code=404, detail="Item not found")
    
    item.is_active = not item.is_active
    item.updated_at = datetime.utcnow()
    db.commit()
    
    return {"status": "ok", "is_active": item.is_active}